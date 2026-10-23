"""Revisão automática de consistência da minuta (primeira etapa do agente-revisor).

Uso:
    python revisar_minuta.py --pasta "<pasta do caso>"           (revisa a minuta mais recente)
    python revisar_minuta.py --pasta "<pasta do caso>" --minuta minuta_v2.docx

Confere, sem julgar o mérito:
    - marcadores não resolvidos, [ERRO], [VERIFICAR], [COMPARAÇÃO VISUAL...] e lacunas (___, XXX, [preencher]);
    - seções obrigatórias e Resumo em documento extenso;
    - figuras: citadas no texto, numeradas em ordem, com legenda, conferidas, com data da imagem;
    - anexos: numeração sequencial pela ordem de citação e lista completa no fim;
    - áreas em hectares citadas no texto que não batem com nenhuma área calculada no caso;
    - coordenadas em grau decimal no texto (o padrão é GMS);
    - o pedido da demanda, para a conferência humana de que o texto responde ao que foi pedido.

Saída: 05_Documento/revisao_vN.md (N = versão da minuta). Código 2 se houver pendência bloqueante.
"""

import argparse
import json
import re
import sys
import unicodedata
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent))
import caso  # noqa: E402

PALAVRAS_RESUMO = 2500


def sem_acento(t):
    return unicodedata.normalize("NFKD", t).encode("ascii", "ignore").decode().lower()


def br_para_float(t):
    return float(t.replace(".", "").replace(",", "."))


def texto_docx(caminho):
    doc = Document(caminho)
    pars = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for linha in t.rows:
            pars.append(" | ".join(c.text for c in linha.cells))
    imagens = sum(1 for p in doc.paragraphs for r in p.runs if r._element.findall(".//{*}drawing"))
    return pars, imagens


def revisar(pasta, minuta):
    ficha = json.loads((pasta / "00_Ficha" / "ficha.json").read_text(encoding="utf-8"))
    pars, n_imagens = texto_docx(minuta)
    tudo = "\n".join(pars)
    bloqueios, atencao, ok = [], [], []

    # 1. pendências de texto
    for padrao, nome in [(r"\[\[[^\]]*\]\]", "marcador não resolvido"), (r"\[ERRO[^\]]*\]", "erro de montagem"),
                         (r"\[VERIFICAR[^\]]*\]", "trecho marcado para verificar"),
                         (r"\[COMPARA[ÇC][ÃA]O VISUAL[^\]]*\]", "comparação visual do NDVI não escrita"), (r"_{3,}|\bX{3,}\b|\[preencher[^\]]*\]|\[[^\]]*N[ÃA]O INFORMAD[OA]\]", "lacuna a preencher")]:
        achados = re.findall(padrao, tudo, flags=re.I)
        if achados:
            bloqueios.append(f"{len(achados)} {nome}(s): {', '.join(sorted(set(achados))[:5])}")
    if not bloqueios:
        ok.append("Nenhum marcador, erro ou lacuna pendente.")

    # 2. seções
    titulos = [sem_acento(p) for p in pars if re.match(r"^\d+\.\s+[A-ZÁÉÍÓÚÂÊÔÃÕÇ ]{4,}$", p.strip())]
    def tem(*nomes):
        return any(any(n in t for n in nomes) for t in titulos)
    for nomes, rotulo in [(("objetivo",), "Objetivo"), (("metodologia",), "Metodologia"), (("analise",), "Análises"),
                          (("resultado", "conclus"), "Resultados e/ou Conclusão")]:
        (ok if tem(*nomes) else bloqueios).append(f"Seção {rotulo}: {'presente' if tem(*nomes) else 'AUSENTE'}.")
    palavras = len(tudo.split())
    if palavras > PALAVRAS_RESUMO and not tem("resumo"):
        atencao.append(f"Documento extenso ({palavras} palavras) sem seção Resumo.")

    # 3. figuras
    citadas = [int(n) for n in re.findall(r"\bFigura\s+(\d+)\b", tudo)]
    legendas = sorted({int(n) for n in re.findall(r"^Figura\s+(\d+)\s+[–-]", tudo, flags=re.M)})
    for n in legendas:
        mencoes = [p for p in pars if re.search(rf"\bFigura\s+{n}\b", p) and not re.match(rf"^Figura\s+{n}\s+[–-]", p)]
        if not mencoes:
            atencao.append(f"Figura {n} inserida, mas não citada no texto.")
    if legendas and legendas != list(range(1, len(legendas) + 1)):
        bloqueios.append(f"Numeração das figuras fora de sequência: {legendas}.")
    for n in sorted(set(citadas) - set(legendas)):
        bloqueios.append(f"Texto cita a Figura {n}, que não está no documento.")
    if n_imagens < len(legendas):
        bloqueios.append(f"{len(legendas)} legendas de figura, mas só {n_imagens} imagens no documento.")
    for f in ficha.get("figuras", []):
        n = int(f["figura"].split()[-1])
        if n in legendas:
            if not f.get("conferido"):
                atencao.append(f"Figura {n} não marcada como conferida (padrão cartográfico).")
            jsons = sorted((pasta / "04_Figuras").glob(f"Figura_{n:02d}_*.json"))
            fmeta = json.loads(jsons[0].read_text(encoding="utf-8")) if jsons else f
            texto_fig = f"{fmeta.get('legenda_figura', '')} {fmeta.get('fonte', '')}"
            usa_imagem = any(i.get("camada") for i in fmeta.get("imagens", []))
            if usa_imagem and not re.search(r"\d{2}/\d{2}/\d{4}|\b(19|20)\d{2}\b", texto_fig):
                atencao.append(f"Figura {n} usa imagem de satélite, mas título e fonte não trazem a data da imagem.")
    if legendas:
        ok.append(f"Figuras no documento: {legendas}.")

    # 3b. tabelas numeradas (ABNT: identificação acima, fonte abaixo)
    tabelas = [int(n) for n in re.findall(r"^Tabela\s+(\d+)\s+[–-]", tudo, flags=re.M)]
    if tabelas and tabelas != list(range(1, len(tabelas) + 1)):
        bloqueios.append(f"Numeração das tabelas fora de sequência: {tabelas}.")
    for n in tabelas:
        if not any(re.search(rf"\bTabela\s+{n}\b", p) and not re.match(rf"^Tabela\s+{n}\s+[–-]", p) for p in pars):
            atencao.append(f"Tabela {n} não citada no texto.")
    n_fontes = len([p for p in pars if p.startswith("Fonte:")])
    if n_fontes < len(legendas) + len(tabelas):
        atencao.append(f"{len(legendas) + len(tabelas)} figuras/tabelas e só {n_fontes} indicações de fonte (ABNT pede fonte abaixo de cada uma).")

    # 4. anexos
    anexos_json = pasta / "05_Documento" / "anexos.json"
    anexos = json.loads(anexos_json.read_text(encoding="utf-8")) if anexos_json.exists() else []
    citados = []
    for n in re.findall(r"\bAnexo\s+(\d+)\b", "\n".join(p for p in pars if not re.match(r"^Anexo\s+\d+\s+[–-]", p))):
        if int(n) not in citados:
            citados.append(int(n))
    if citados != list(range(1, len(citados) + 1)):
        bloqueios.append(f"Anexos citados fora da ordem de numeração: {citados}.")
    lista = [int(n) for n in re.findall(r"^Anexo\s+(\d+)\s+[–-]", tudo, flags=re.M)]
    if sorted(set(citados)) != sorted(set(lista)):
        bloqueios.append(f"Anexos citados {sorted(set(citados))} diferentes da Relação de Anexos {sorted(set(lista))}.")
    elif lista:
        ok.append(f"Anexos 1 a {max(lista)} citados em ordem e listados.")
    pasta_anexos = pasta / "06_Anexos"
    faltando = [a["arquivo_sugerido"] for a in anexos if not list(pasta_anexos.glob(f"Anexo_{a['numero']}_*"))]
    if faltando:
        atencao.append(f"Arquivos de anexo ainda não preparados em 06_Anexos: {', '.join(faltando)} (etapa de anexos).")

    # 5. áreas citadas x áreas calculadas
    calculadas = [round(a["area_ha"], 2) for a in ficha.get("areas", []) if isinstance(a, dict) and a.get("area_ha")]
    for txt in sorted(set(re.findall(r"(\d{1,3}(?:\.\d{3})*,\d+)\s*(?:ha|hectares)\b", tudo))):
        v = round(br_para_float(txt), 2)
        if calculadas and not any(abs(v - c) <= 0.01 for c in calculadas):
            atencao.append(f"Área {txt} ha citada no texto não bate com nenhuma área calculada no caso {calculadas}. Confira a origem.")

    # 6. coordenadas em grau decimal
    decimais = re.findall(r"-\d{1,2}[.,]\d{4,}", tudo)
    if decimais:
        atencao.append(f"Coordenadas em grau decimal no texto ({', '.join(decimais[:3])}); o padrão é GMS.")

    return {"bloqueios": bloqueios, "atencao": atencao, "ok": ok, "palavras": palavras,
            "pedido": ficha.get("pedido"), "tipo": ficha.get("tipo_documento")}


def main(argv=None):
    ap = argparse.ArgumentParser(description="Revisão automática de consistência da minuta.")
    ap.add_argument("--pasta", required=True)
    ap.add_argument("--minuta")
    a = ap.parse_args(argv)
    pasta = Path(a.pasta)
    pasta_doc = pasta / "05_Documento"
    if a.minuta:
        minuta = pasta_doc / a.minuta
    else:
        versoes = sorted(pasta_doc.glob("minuta_v*.docx"), key=lambda f: int(re.search(r"v(\d+)", f.stem).group(1)))
        if not versoes:
            sys.exit("Nenhuma minuta em 05_Documento.")
        minuta = versoes[-1]
    v = re.search(r"v(\d+)", minuta.stem).group(1)
    r = revisar(pasta, minuta)
    linhas = [f"# Revisão automática – {minuta.name}", "",
              f"Documento: {r['tipo']} · {r['palavras']} palavras", "",
              "## Pendências que impedem seguir", ""] + ([f"- {x}" for x in r["bloqueios"]] or ["- Nenhuma."])
    linhas += ["", "## Pontos de atenção", ""] + ([f"- {x}" for x in r["atencao"]] or ["- Nenhum."])
    linhas += ["", "## Conferido", ""] + [f"- {x}" for x in r["ok"]]
    linhas += ["", "## Para a revisão humana e do agente-revisor", "",
               f"- O texto responde ao pedido da demanda? Pedido: \"{r['pedido'] or '(não registrado na ficha)'}\"",
               "- As conclusões decorrem das análises e dos resultados apresentados?",
               "- A linguagem é clara para quem não é especialista (advogados, chefias, procuradores)?",
               "- Os mapas atendem ao padrão cartográfico (objetivo, primeiro plano, rótulos, linhas, imagem)?"]
    saida = pasta_doc / f"revisao_v{v}.md"
    saida.write_text("\n".join(linhas) + "\n", encoding="utf-8")
    caso.registrar_linha(pasta, "revisao", f"Revisão automática da minuta v{v}: {len(r['bloqueios'])} pendência(s), {len(r['atencao'])} ponto(s) de atenção.")
    print("\n".join(linhas))
    sys.exit(2 if r["bloqueios"] else 0)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
