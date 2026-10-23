"""Monta a minuta em Word a partir do conteúdo escrito em texto com marcadores.

Uso:
    python minuta.py --pasta "<pasta do caso>"            (lê 05_Documento/conteudo.md)
    python minuta.py --pasta "<pasta do caso>" --conteudo outro.md

Requer: pip install python-docx

Marcadores no conteúdo (cada um resolvido a partir dos dados do caso, sem digitação de números):
    [[FIGURA:1]]                          sozinho na linha: insere 04_Figuras/Figura_01_*.png com a legenda do JSON
    [[ANEXO:chave|Descrição do anexo]]    vira "Anexo N", numerado pela ordem da primeira citação;
                                          nas citações seguintes basta [[ANEXO:chave]]
    [[AREA:camada]]                       área da ficha, ex.: "78,63 ha"
    [[PERIMETRO:camada]]                  perímetro da ficha, ex.: "3.550,29 m"
    [[TABELA_COORDENADAS:camada]]         sozinho na linha: tabela de vértices em GMS (SIRGAS 2000) do caso.gpkg
    [[NUP]] [[TIPO]] [[DEMANDANTE]] [[DOCUMENTO_DEMANDA]]   dados da ficha

Formatação aceita: ## Seção, ### Subseção, parágrafos, **negrito**, *itálico*, listas com "- " e "1. ",
tabelas simples com "|". A seção "Relação de Anexos" é gerada sozinha no fim (não precisa escrever).

Saídas em 05_Documento/:
    minuta_vN.docx   nova versão a cada execução (nunca sobrescreve)
    anexos.json      lista numerada de anexos, usada para nomear os arquivos
"""

import argparse
import json
import re
import sqlite3
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
import caso  # noqa: E402

FONTE = "Calibri"
TAM = 12
SECOES_OBRIGATORIAS = ["objetivo", "metodologia", "analises", ("resultados", "conclusao")]
RE_ANEXO = re.compile(r"\[\[ANEXO:([A-Za-z0-9_\-]+)(?:\|([^\]]+))?\]\]")
RE_MARCADOR = re.compile(r"\[\[([A-Z_]+)(?::([^\]]+))?\]\]")


def agora():
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def sem_acento(t):
    return unicodedata.normalize("NFKD", t).encode("ascii", "ignore").decode().lower()


def slug(t, n=40):
    return re.sub(r"[^A-Za-z0-9]+", "_", unicodedata.normalize("NFKD", t).encode("ascii", "ignore").decode()).strip("_")[:n]


def br(valor, casas=2):
    return f"{valor:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def ler_conteudo(caminho):
    texto = Path(caminho).read_text(encoding="utf-8")
    meta = {}
    if texto.startswith("---"):
        fim = texto.find("\n---", 3)
        for linha in texto[3:fim].strip().splitlines():
            if ":" in linha:
                k, v = linha.split(":", 1)
                meta[k.strip()] = v.strip()
        texto = texto[fim + 4:]
    return meta, texto


def numerar_anexos(texto):
    """Numera os anexos pela ordem da primeira citação. Devolve lista e erros."""
    ordem, descricoes, erros = [], {}, []
    for m in RE_ANEXO.finditer(texto):
        chave, desc = m.group(1), m.group(2)
        if chave not in ordem:
            ordem.append(chave)
            if not desc:
                erros.append(f"Primeira citação do anexo '{chave}' sem descrição: use [[ANEXO:{chave}|Descrição]].")
        if desc and chave not in descricoes:
            descricoes[chave] = desc.strip()
    anexos = [{"numero": i, "chave": c, "descricao": descricoes.get(c, c),
               "arquivo_sugerido": f"Anexo_{i}_" + "_".join(w.capitalize() for w in re.split(r"[_\-]+", c) if w)} for i, c in enumerate(ordem, start=1)]
    return anexos, erros


# ---------- formatação ----------

def estilo_base(doc):
    st = doc.styles["Normal"]
    st.font.name = FONTE
    st.font.size = Pt(TAM)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
    pf = st.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(2.5), Cm(2)
        s.left_margin, s.right_margin = Cm(3), Cm(2)


def texto_com_formatacao(par, texto):
    """Escreve o texto no parágrafo tratando **negrito** e *itálico*."""
    for parte in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            par.add_run(parte[2:-2]).bold = True
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            par.add_run(parte[1:-1]).italic = True
        else:
            par.add_run(parte)


def titulo_secao(doc, texto, nivel, numero):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if nivel == 2 else 8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{numero} {texto.upper() if nivel == 2 else texto}")
    r.bold = True
    r.font.size = Pt(TAM)


def borda_tabela(tabela):
    tbl = tabela._tbl
    pr = tbl.tblPr
    bordas = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "808080")
        bordas.append(el)
    pr.append(bordas)


def tabela(doc, linhas, cabecalho=True, tamanho=10, manter_junta=20):
    """Tabela com cabeçalho repetido em cada página; até `manter_junta` linhas, não quebra entre páginas."""
    t = doc.add_table(rows=len(linhas), cols=len(linhas[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borda_tabela(t)
    junta = len(linhas) <= manter_junta
    for i, linha in enumerate(linhas):
        tr_pr = t.rows[i]._tr.get_or_add_trPr()
        nao_quebra = OxmlElement("w:cantSplit")
        tr_pr.append(nao_quebra)
        if cabecalho and i == 0:
            repete = OxmlElement("w:tblHeader")
            tr_pr.append(repete)
        for j, valor in enumerate(linha):
            cel = t.cell(i, j)
            cel.text = ""
            p = cel.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            if junta and i < len(linhas) - 1:
                p.paragraph_format.keep_with_next = True
            r = p.add_run(str(valor))
            r.font.size = Pt(tamanho)
            if cabecalho and i == 0:
                r.bold = True
                sombra = OxmlElement("w:shd")
                sombra.set(qn("w:val"), "clear")
                sombra.set(qn("w:fill"), "EDEDED")
                cel._tc.get_or_add_tcPr().append(sombra)
    return t


def titulo_ilustracao(doc, texto):
    """Identificação acima da figura ou tabela (ABNT NBR 14724)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(texto)
    r.font.size = Pt(10)
    r.bold = True


def fonte_ilustracao(doc, texto):
    """Fonte abaixo da figura ou tabela (ABNT NBR 14724)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(texto if texto.lower().startswith("fonte") else f"Fonte: {texto}")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(70, 70, 70)


# ---------- dados do caso ----------

def figura(pasta, numero):
    jsons = sorted((pasta / "04_Figuras").glob(f"Figura_{int(numero):02d}_*.json"))
    if not jsons:
        raise ValueError(f"Figura {numero} não existe em 04_Figuras (gere com figura_mapa.py).")
    meta = json.loads(jsons[0].read_text(encoding="utf-8"))
    return meta, jsons[0].with_name(meta["arquivo"])


def area(ficha, camada, campo):
    for a in ficha.get("areas", []):
        if isinstance(a, dict) and a.get("camada") == camada and a.get(campo) is not None:
            return a[campo]
    raise ValueError(f"Sem {campo} da camada '{camada}' na ficha (rode projeto_qgis.py).")


def coordenadas(pasta, camada):
    gpkg = pasta / "03_Geo" / "vetores_produzidos" / "caso.gpkg"
    con = sqlite3.connect(gpkg)
    try:
        linhas = con.execute(f'SELECT vertice, lat_gms, lon_gms FROM "{camada}" ORDER BY ordem').fetchall()
    except sqlite3.OperationalError as e:
        raise ValueError(f"Não li a camada '{camada}' em caso.gpkg: {e}")
    finally:
        con.close()
    return [["Vértice", "Latitude", "Longitude"]] + [list(l) for l in linhas]


# ---------- montagem ----------

def montar(pasta, conteudo_md):
    ficha_json = pasta / "00_Ficha" / "ficha.json"
    ficha = json.loads(ficha_json.read_text(encoding="utf-8"))
    meta, texto = ler_conteudo(conteudo_md)
    anexos, erros = numerar_anexos(texto)
    num_anexo = {a["chave"]: a["numero"] for a in anexos}
    alertas, figuras_usadas = [], []

    def inline(t):
        t = RE_ANEXO.sub(lambda m: f"Anexo {num_anexo[m.group(1)]}", t)

        def troca(m):
            nome, arg = m.group(1), m.group(2)
            try:
                if nome == "AREA":
                    return f"{br(area(ficha, arg, 'area_ha'))} ha"
                if nome == "PERIMETRO":
                    return f"{br(area(ficha, arg, 'perimetro_m'))} m"
                if nome == "NUP":
                    return ficha["nup"]
                if nome == "TIPO":
                    return ficha["tipo_documento"]
                if nome == "DEMANDANTE":
                    return ficha.get("setor_demandante") or "[DEMANDANTE NÃO INFORMADO]"
                if nome == "DOCUMENTO_DEMANDA":
                    return ficha.get("documento_demanda_sei") or "[Nº SEI NÃO INFORMADO]"
            except ValueError as e:
                erros.append(str(e))
                return f"[ERRO: {m.group(0)}]"
            return m.group(0)
        return RE_MARCADOR.sub(troca, t)

    doc = Document()
    estilo_base(doc)

    # identificação curta (o cabeçalho oficial é o do SEI)
    for rotulo, valor in [("Processo", ficha["nup"]), ("Assunto", meta.get("assunto") or ficha.get("pedido") or "")]:
        if valor:
            p = doc.add_paragraph()
            p.add_run(f"{rotulo}: ").bold = True
            p.add_run(valor)

    sec, sub, n_tab = 0, 0, 0
    titulo_tabela_seguinte, fonte_pendente = None, False
    secoes_vistas = []
    bloco_lista = None
    linhas = texto.splitlines()
    i = 0
    while i < len(linhas):
        l = linhas[i].rstrip()
        s = l.strip()
        if not s:
            bloco_lista = None
            i += 1
            continue
        if s.startswith("## "):
            nome = s[3:].strip()
            if sem_acento(nome).startswith("relacao de anexos"):
                i += 1
                while i < len(linhas) and not linhas[i].startswith("## "):
                    i += 1
                continue
            sec, sub = sec + 1, 0
            secoes_vistas.append(sem_acento(nome))
            titulo_secao(doc, nome, 2, f"{sec}.")
        elif s.startswith("### "):
            sub += 1
            titulo_secao(doc, s[4:].strip(), 3, f"{sec}.{sub}.")
        elif s.startswith("# "):
            pass  # o título do documento é o do SEI
        elif re.fullmatch(r"\[\[FIGURA:(\d+)\]\]", s):
            n = re.fullmatch(r"\[\[FIGURA:(\d+)\]\]", s).group(1)
            try:
                fmeta, png = figura(pasta, n)
                titulo_ilustracao(doc, fmeta["legenda_figura"])
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_after = Pt(0)
                p.add_run().add_picture(str(png), width=Cm(16))
                fonte_ilustracao(doc, fmeta.get("fonte") or "elaborado pelo autor. Sistema de referência SIRGAS 2000.")
                figuras_usadas.append(int(n))
                if not fmeta.get("conferido"):
                    alertas.append(f"Figura {n} ainda não foi marcada como conferida.")
            except ValueError as e:
                erros.append(str(e))
                doc.add_paragraph(f"[ERRO: figura {n} não encontrada]")
        elif re.fullmatch(r"\[\[TABELA_COORDENADAS:([^\]|]+)(?:\|([^\]]+))?\]\]", s):
            m = re.fullmatch(r"\[\[TABELA_COORDENADAS:([^\]|]+)(?:\|([^\]]+))?\]\]", s)
            camada, titulo = m.group(1), m.group(2) or "Coordenadas dos vértices"
            try:
                linhas_tab = coordenadas(pasta, camada)
                n_tab += 1
                titulo_ilustracao(doc, f"Tabela {n_tab} – {titulo}")
                tabela(doc, linhas_tab)
                fonte_ilustracao(doc, "elaborado pelo autor. Sistema de referência SIRGAS 2000; coordenadas em graus-minutos-segundos.")
            except ValueError as e:
                erros.append(str(e))
        elif s.lower().startswith("tabela:") and i + 1 < len(linhas) and linhas[i + 1].strip().startswith("|"):
            titulo_tabela_seguinte = inline(s.split(":", 1)[1].strip())
        elif s.lower().startswith("fonte:") and fonte_pendente:
            fonte_ilustracao(doc, inline(s))
            fonte_pendente = False
        elif s.startswith("|"):
            if titulo_tabela_seguinte:
                n_tab += 1
                titulo_ilustracao(doc, f"Tabela {n_tab} – {titulo_tabela_seguinte}")
                titulo_tabela_seguinte = None
            bloco = []
            while i < len(linhas) and linhas[i].strip().startswith("|"):
                celulas = [inline(c.strip()) for c in linhas[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in celulas):
                    bloco.append(celulas)
                i += 1
            tabela(doc, bloco)
            fonte_pendente = True
            continue
        elif re.match(r"^(- |\* )", s):
            p = doc.add_paragraph(style="List Bullet")
            texto_com_formatacao(p, inline(s[2:]))
        elif re.match(r"^\d+\. ", s):
            p = doc.add_paragraph(style="List Number")
            texto_com_formatacao(p, inline(re.sub(r"^\d+\. ", "", s)))
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            texto_com_formatacao(p, inline(s))
        i += 1

    if anexos:
        sec += 1
        titulo_secao(doc, "Relação de Anexos", 2, f"{sec}.")
        for a in anexos:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.add_run(f"Anexo {a['numero']} – ").bold = True
            p.add_run(a["descricao"])

    for obrig in SECOES_OBRIGATORIAS:
        nomes = obrig if isinstance(obrig, tuple) else (obrig,)
        if not any(any(v.startswith(n) for n in nomes) for v in secoes_vistas):
            alertas.append(f"Seção obrigatória ausente: {' ou '.join(nomes)}.")

    pasta_doc = pasta / "05_Documento"
    pasta_doc.mkdir(parents=True, exist_ok=True)
    versoes = [int(m.group(1)) for f in pasta_doc.glob("minuta_v*.docx") if (m := re.match(r"minuta_v(\d+)\.docx", f.name))]
    versao = max(versoes, default=0) + 1
    saida = pasta_doc / f"minuta_v{versao}.docx"
    doc.save(saida)
    (pasta_doc / "anexos.json").write_text(json.dumps(anexos, ensure_ascii=False, indent=2), encoding="utf-8")

    ficha["anexos"] = [f"Anexo {a['numero']} – {a['descricao']}" for a in anexos]
    ficha_json.write_text(json.dumps(ficha, ensure_ascii=False, indent=2), encoding="utf-8")
    caso.gravar_ficha_md(pasta, ficha)
    caso.registrar_linha(pasta, "redacao", f"Minuta v{versao} gerada ({len(anexos)} anexos, figuras {figuras_usadas or 'nenhuma'}).")

    return {"minuta": str(saida), "versao": versao, "anexos": anexos, "figuras": figuras_usadas,
            "erros": erros, "alertas": alertas, "em": agora()}


def main(argv=None):
    ap = argparse.ArgumentParser(description="Monta a minuta em Word a partir do conteúdo com marcadores.")
    ap.add_argument("--pasta", required=True)
    ap.add_argument("--conteudo", help="Padrão: 05_Documento/conteudo.md")
    a = ap.parse_args(argv)
    pasta = Path(a.pasta)
    conteudo = Path(a.conteudo) if a.conteudo else pasta / "05_Documento" / "conteudo.md"
    if not conteudo.exists():
        sys.exit(f"Conteúdo não encontrado: {conteudo}")
    r = montar(pasta, conteudo)
    print(json.dumps(r, ensure_ascii=False, indent=2))
    sys.exit(2 if r["erros"] else 0)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
